"""Simple Markdown-style extraction tests. No Docker or network."""

import io

import pytest
from docx import Document as DocxDocument

from app.config import Settings
from app.exceptions import UnsupportedMediaTypeError
from app.services.text_extraction_service import (
    ExtractionError,
    TextExtractionService,
    _extract_sync,
)


def _make_pdf(page_texts: list[str]) -> bytes:
    page_nums = [4 + 2 * i for i in range(len(page_texts))]
    content_nums = [5 + 2 * i for i in range(len(page_texts))]
    objects: dict[int, bytes] = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        2: (
            f"<< /Type /Pages /Kids "
            f"[{' '.join(f'{number} 0 R' for number in page_nums)}] "
            f"/Count {len(page_texts)} >>"
        ).encode(),
        3: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    }
    for index, value in enumerate(page_texts):
        escaped = value.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
        stream = f"BT /F1 12 Tf 20 150 Td ({escaped}) Tj ET".encode()
        objects[content_nums[index]] = (
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"
        )
        objects[page_nums[index]] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 200] "
            f"/Contents {content_nums[index]} 0 R "
            f"/Resources << /Font << /F1 3 0 R >> >> >>"
        ).encode()
    output = bytearray(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}
    for number in sorted(objects):
        offsets[number] = len(output)
        output += f"{number} 0 obj\n".encode() + objects[number] + b"\nendobj\n"
    xref = len(output)
    highest = max(objects)
    output += f"xref\n0 {highest + 1}\n".encode() + b"0000000000 65535 f \n"
    for number in range(1, highest + 1):
        output += f"{offsets[number]:010d} 00000 n \n".encode()
    output += (
        f"trailer\n<< /Size {highest + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
    ).encode()
    return bytes(output)


def test_txt_groups_lines_into_a_normalised_paragraph() -> None:
    document = _extract_sync(b"line one\r\nline two", "notes.txt")
    assert [block.content for block in document.blocks] == ["line one line two"]


def test_markdown_carries_the_section_path_and_link() -> None:
    document = _extract_sync(
        b"# Leave\n\n## Carry over\n\nRead the [policy](https://example.com/leave).",
        "readme.md",
    )
    block = document.blocks[-1]
    assert block.section_path == "Leave > Carry over"
    assert "[policy](https://example.com/leave)" in block.content


def test_csv_becomes_one_complete_table() -> None:
    document = _extract_sync(b"name,dept\nAlice,Finance\nBob,HR\n", "people.csv")
    assert len(document.blocks) == 1
    assert document.blocks[0].block_type == "table"
    assert document.blocks[0].content.splitlines() == [
        "name | dept",
        "Alice | Finance",
        "Bob | HR",
    ]


def test_pdf_text_becomes_paragraph_blocks() -> None:
    document = _extract_sync(_make_pdf(["alpha", "beta"]), "doc.pdf")
    assert [block.content for block in document.blocks] == ["alpha", "beta"]


def test_docx_heading_and_paragraph_become_simple_blocks() -> None:
    doc = DocxDocument()
    doc.add_paragraph("Leave Policy", style="Heading 1")
    doc.add_paragraph("Annual leave applies.")
    buffer = io.BytesIO()
    doc.save(buffer)
    document = _extract_sync(buffer.getvalue(), "policy.docx")
    assert [block.block_type for block in document.blocks] == ["paragraph"]
    assert document.blocks[0].section_path == "Leave Policy"


@pytest.mark.parametrize("name", ["broken.pdf", "broken.docx"])
def test_corrupt_structured_file_raises(name: str) -> None:
    with pytest.raises(ExtractionError):
        _extract_sync(b"not a real document", name)


def test_unsupported_extension_is_rejected() -> None:
    with pytest.raises(UnsupportedMediaTypeError):
        _extract_sync(b"...", "archive.zip")


def test_supported_extensions_match_upload_allowlist() -> None:
    from app.services.text_extraction_service import _HANDLERS

    assert set(_HANDLERS) == set(Settings().allowed_upload_extensions)


async def test_service_returns_extracted_markdown() -> None:
    document = await TextExtractionService().extract(b"hello", "a.txt")
    assert document.blocks[0].content == "hello"
    assert not document.is_empty


def test_markdown_preserves_sections_lists_tables_and_code() -> None:
    document = _extract_sync(
        b"# Policy\n\nIntro.\n\n1. Parent\n   - Child\n\n"
        b"| Key | Value |\n| --- | --- |\n| A | B |\n\n"
        b"```py\ndef run(): pass\n```",
        "policy.md",
    )
    assert {block.block_type for block in document.blocks} == {
        "paragraph",
        "list",
        "table",
        "code",
    }
    assert all(block.section_path == "Policy" for block in document.blocks)
    table = next(block for block in document.blocks if block.block_type == "table")
    assert table.content.splitlines() == ["Key | Value", "A | B"]
